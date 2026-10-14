"""Auditoría exacta de sitios MER contra catálogo oficial por país."""
from io import BytesIO
from pathlib import Path
from difflib import get_close_matches
import json
import re
import unicodedata
import pandas as pd
from openpyxl import load_workbook
from auditoria_tablas import DICCIONARIO

RUTA = Path(__file__).with_name('diccionario_sitios_mer.json')
CATALOGO = json.loads(RUTA.read_text(encoding='utf-8'))['sitios']
CAMPOS = ['Archivo', 'Ruta', 'Hoja', 'Fila Excel', 'País reportado', 'Departamento reportado', 'Columna de servicio', 'Sitio encontrado', 'Sitio oficial sugerido', 'País oficial', 'Estado', 'Detalle']
SITIO_COLUMNAS = ['Servicio de Salud (Unidad de Tratamiento)', 'Servicio de Salud', 'Nombre del Servicio de Salud', 'Nombre del establecimiento', 'Establecimiento de Salud', 'Sitio', 'Nombre del sitio']
PAIS_COLUMNAS = ['País', 'Pais', 'País de reporte']
DEPTO_COLUMNAS = ['Departamento', 'DEPARTAMENTO']

def base(s):
    return ''.join(c for c in unicodedata.normalize('NFKD', str(s).casefold()) if not unicodedata.combining(c))

def suave(s):
    return re.sub(r'\s+', ' ', base(s).strip())

def detectar(headers, opciones):
    for opcion in opciones:
        for i, encabezado in enumerate(headers):
            if isinstance(encabezado, str) and suave(encabezado) == suave(opcion):
                return i
    return None

def auditar_sitios_excel(entradas):
    filas = []
    por_pais = {}
    for r in CATALOGO:
        por_pais.setdefault(suave(r['pais']), []).append(r)
    for archivo, contenido, ruta in entradas:
        try:
            wb = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
        except Exception as exc:
            filas.append([archivo, ruta, '', '', '', '', '', '', '', '', 'ERROR LECTURA', str(exc)[:250]])
            continue
        try:
            for ws in wb.worksheets:
                # Solo las hojas de indicadores presentes en el formato MER oficial.
                if ws.title not in DICCIONARIO:
                    continue
                # Identificar encabezados en las primeras 18 filas.
                encabezado = None
                for numero, fila in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row,18), values_only=True), 1):
                    idx = detectar(fila, SITIO_COLUMNAS)
                    if idx is not None:
                        encabezado = (numero, list(fila), idx)
                        break
                if encabezado is None:
                    continue
                inicio, columnas, ix = encabezado
                ip = detectar(columnas, PAIS_COLUMNAS)
                idep = detectar(columnas, DEPTO_COLUMNAS)
                vistos = set()
                for num, row in enumerate(ws.iter_rows(min_row=inicio+1, values_only=True),inicio+1):
                    if ix >= len(row):
                        continue
                    valor = row[ix]
                    if valor is None or not str(valor).strip():
                        continue
                    sitio = str(valor)
                    pais = str(row[ip]).strip() if ip is not None and ip < len(row) and row[ip] is not None else ''
                    depto = str(row[idep]).strip() if idep is not None and idep < len(row) and row[idep] is not None else ''
                    clave = (sitio, suave(pais))
                    if clave in vistos:
                        continue
                    vistos.add(clave)
                    candidatos = por_pais.get(suave(pais), []) if pais else CATALOGO
                    iguales = [r for r in candidatos if r['sitio'] == sitio]
                    if iguales:
                        estado, sugerido, pais_oficial, detalle = 'CORRECTO',sitio,iguales[0]['pais'],'Coincidencia exacta con el catálogo oficial.'
                    else:
                        approx = [r for r in candidatos if suave(r['sitio']) == suave(sitio)]
                        if not approx:
                            nombres = [r['sitio'] for r in candidatos]
                            posibles = get_close_matches(sitio,nombres,n=1,cutoff=0.59)
                            approx = [r for r in candidatos if r['sitio'] == posibles[0]] if posibles else []
                        sug = approx[0] if approx else None
                        sugerido = sug['sitio'] if sug else ''
                        pais_oficial = sug['pais'] if sug else ''
                        if sug and suave(sitio) == suave(sugerido):
                            estado,detalle='DIFERENCIA DE CARACTERES','Diferencia de espacios, tildes o mayúsculas; se requiere coincidencia exacta.'
                        elif sug:
                            estado,detalle='NOMBRE NO OFICIAL','Nombre diferente al catálogo oficial; verificar propuesta (no se corrige automáticamente).'
                        else:
                            estado,detalle='SITIO NO ENCONTRADO','No figura en el catálogo de sitios oficiales para el país identificado.'
                    filas.append([archivo,ruta,ws.title,num,pais,depto,columnas[ix],sitio,sugerido,pais_oficial,estado,detalle])
        finally:
            wb.close()
    return pd.DataFrame(filas,columns=CAMPOS)
