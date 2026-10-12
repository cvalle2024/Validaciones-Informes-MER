"""Auditoría de nombres de tablas estructuradas y hojas de Excel MER."""
from collections import Counter, defaultdict
from io import BytesIO
from openpyxl import load_workbook
import pandas as pd

CAMPOS = ["Archivo", "Ruta", "Hoja", "Tabla esperada", "Tabla encontrada", "Estado", "Detalle"]


def auditar_tablas_excel(entradas, referencia=None):
    """Compara nombres exactos por hoja frente al archivo de referencia o al consenso.

    entradas: [(filename, file_bytes, relative_path), ...].
    referencia: relative_path de un archivo elegido, o None (consenso de archivos).
    Nunca ignora silenciosamente libros/hojas que no se pueden abrir.
    """
    libros = []
    hallazgos = []
    for nombre, contenido, ruta in entradas:
        try:
            wb = load_workbook(BytesIO(contenido), read_only=False, data_only=False)
            inventario = {ws.title: set(ws.tables.keys()) for ws in wb.worksheets}
            wb.close()
            libros.append((nombre, ruta, inventario))
        except Exception as exc:
            hallazgos.append({"Archivo": nombre, "Ruta": ruta, "Hoja": "", "Tabla esperada": "", "Tabla encontrada": "", "Estado": "ERROR LECTURA", "Detalle": str(exc)[:350]})

    if not libros:
        return pd.DataFrame(hallazgos, columns=CAMPOS), pd.DataFrame(columns=["Hoja", "Tabla esperada", "Archivos que la contienen", "Origen"])

    # Tomar como referencia los nombres de tablas definidos en un archivo, o
    # la combinación más frecuente para cada hoja presente en >=2 archivos.
    selected = next((d for n, r, d in libros if r == referencia), None)
    sheet_presence = Counter(hoja for _, _, inv in libros for hoja in inv)
    profiles = defaultdict(Counter)
    for _, _, inv in libros:
        for sheet, names in inv.items():
            profiles[sheet][tuple(sorted(names))] += 1

    # Para detectar ausencias, tomar solo las hojas recurrentes, salvo con archivo patrón.
    if selected is not None:
        canonical = {s: set(names) for s, names in selected.items()}
        origen = "Archivo de referencia"
    else:
        canonical = {}
        for sheet, counts in profiles.items():
            if sheet_presence[sheet] < 2:
                continue
            # Preferir perfiles presentes más frecuentemente; ante empate,
            # priorizar el que tenga más nombres (reduce falsos negativos).
            profile = max(counts.items(), key=lambda kv: (kv[1], len(kv[0]), kv[0]))[0]
            canonical[sheet] = set(profile)
        origen = "Consenso entre archivos"

    for nombre, ruta, inv in libros:
        for sheet in sorted(canonical):
            expected = canonical[sheet]
            if sheet not in inv:
                hallazgos.append({"Archivo": nombre, "Ruta": ruta, "Hoja": sheet,
                                  "Tabla esperada": ", ".join(sorted(expected)), "Tabla encontrada": "",
                                  "Estado": "HOJA FALTANTE", "Detalle": "La hoja esperada no aparece en este libro; revisar si este formato debe contenerla."})
                continue
            observed = inv[sheet]
            missing, extra = sorted(expected - observed), sorted(observed - expected)
            if missing and extra:
                estado = "NOMBRE DIFERENTE"
                detalle = "Faltan: " + ", ".join(missing) + "; aparecen: " + ", ".join(extra)
            elif missing:
                estado = "TABLA FALTANTE"
                detalle = "Tabla(s) estructurada(s) no encontrada(s): " + ", ".join(missing)
            elif extra:
                estado = "TABLA INESPERADA"
                detalle = "Tabla(s) adicionales: " + ", ".join(extra)
            else:
                estado = "CORRECTO"
                detalle = "Nombres exactos coincidentes" if expected else "Hoja sin tablas estructuradas en el patrón"
            hallazgos.append({"Archivo": nombre, "Ruta": ruta, "Hoja": sheet,
                              "Tabla esperada": ", ".join(sorted(expected)),
                              "Tabla encontrada": ", ".join(sorted(observed)), "Estado": estado, "Detalle": detalle})
        for sheet in sorted(set(inv) - set(canonical)):
            hallazgos.append({"Archivo": nombre, "Ruta": ruta, "Hoja": sheet,
                              "Tabla esperada": "Sin referencia", "Tabla encontrada": ", ".join(sorted(inv[sheet])),
                              "Estado": "SIN PATRÓN", "Detalle": "Hoja no incluida en la referencia; no se puede asegurar nombre correcto."})

    patrones = pd.DataFrame([
        {"Hoja": s, "Tabla esperada": ", ".join(sorted(names)),
         "Archivos que la contienen": sheet_presence[s], "Origen": origen}
        for s, names in sorted(canonical.items())
    ], columns=["Hoja", "Tabla esperada", "Archivos que la contienen", "Origen"])
    return pd.DataFrame(hallazgos, columns=CAMPOS), patrones
